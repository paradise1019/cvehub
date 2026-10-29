import io
import os
import unittest

from cvehub import osv_import, period, vex_import
from tests.helpers import TempEnv, make_osv_tar, make_vex_tar, osv_record, usn_record

E20P, E22, E22P, E24 = "Ubuntu:Pro:20.04:LTS", "Ubuntu:22.04:LTS", "Ubuntu:Pro:22.04:LTS", "Ubuntu:24.04:LTS"


def rec(cve, affected, prio, published, details="說明"):
    d = osv_record(cve, affected, prio)
    d["published"] = published
    d["details"] = details
    return d


RECS = [
    # 三個版本都有，20.04 只在 Pro 修好
    rec("CVE-2026-1001", [(E20P, "openssl", "1.1.1f-1ubuntu2.24+esm1", {"openssl": "1.1.1f-1ubuntu2.24+esm1"}),
                          ("Ubuntu:20.04:LTS", "openssl", None, {}),
                          (E22, "openssl", "3.0.2-0ubuntu1.20", {"openssl": "3.0.2-0ubuntu1.20"}),
                          (E24, "openssl", "3.0.13-0ubuntu3.16", {"openssl": "3.0.13-0ubuntu3.16", "libssl3t64": "3.0.13-0ubuntu3.16"})],
        "high", "2026-07-10T00:00:00Z", "openssl 的問題"),
    # kernel：只留 linux 與 linux-hwe-*，雲端變體不列；24.04 尚未修正
    rec("CVE-2026-1002", [(E22, "linux", "5.15.0-160.170", {}), (E22, "linux-hwe-6.8", "6.8.0-90.91~22.04.1", {}),
                          (E22, "linux-aws", "5.15.0-1090.97", {}), (E24, "linux", None, {}),
                          (E24, "linux-firmware", "20240318.git3b128b60-0ubuntu2.20", {})],
        "critical", "2026-08-01T12:00:00Z"),
    # 舊的 HWE 系列（6.5）不是最新、也沒有主機裝 → 不列
    rec("CVE-2026-1006", [(E22, "linux-hwe-6.5", None, {}), (E22, "linux-hwe-6.8", "6.8.0-91.92~22.04.1", {})],
        "high", "2026-08-02T00:00:00Z", "x" * 1500),
    # 22.04 只有 Pro（universe）才有修正
    rec("CVE-2026-1003", [(E22P, "nodejs", "12.22.9~dfsg-1ubuntu3.6+esm1", {})], "medium", "2026-09-30T23:00:00Z"),
    # 區間外、等級太低
    rec("CVE-2026-1004", [(E24, "curl", "8.5.0-2ubuntu10.9", {})], "high", "2026-10-01T00:00:00Z"),
    rec("CVE-2026-1005", [(E24, "curl", "8.5.0-2ubuntu10.9", {})], "low", "2026-08-15T00:00:00Z"),
    # OSV 沒有任何指定版本的套件（只有 26.04）：季報表仍要列出這個 CVE
    rec("CVE-2026-1007", [("Ubuntu:26.04:LTS", "zlib", "1:1.3", {})], "medium", "2026-09-01T00:00:00Z"),
]

USNS = [
    usn_record("USN-1", "2026-09-01", ["CVE-2026-1001"], [(E20P, "openssl", "1.1.1f-1ubuntu2.24+esm1"), (E22, "openssl", "3.0.2-0ubuntu1.20"),
                                                      (E24, "openssl", "3.0.13-0ubuntu3.16")]),
    usn_record("USN-2", "2026-10-05", ["CVE-2026-1002"], [(E22, "linux", "5.15.0-160.170"), (E22, "linux-aws", "5.15.0-1090.97")]),  # 期末後
    usn_record("USN-3", "2026-08-10", ["CVE-2026-1002"], [(E22, "linux-hwe-6.8", "6.8.0-90.91~22.04.1")]),
    usn_record("USN-4", "2026-09-20", ["CVE-2026-1006"], [(E22, "linux-hwe-6.8", "6.8.0-91.92~22.04.1")]),
    usn_record("USN-4b", "2026-09-25", ["CVE-2026-1006"], [(E22, "linux-hwe-6.8", "6.8.0-91.92~22.04.1")]),       # 較晚的重複 USN
    usn_record("USN-5", "2026-09-29", ["CVE-2026-1003"], [(E22P, "nodejs", "12.22.9~dfsg-1ubuntu3.6+esm1")]),
]   # linux-firmware（24.04）沒有 USN → 日期不明

VEX = {
    "CVE-2026-1001": [("fixed", [("openssl", "esm-infra/focal"), ("openssl", "jammy"), ("openssl", "noble")]),
                      ("eol", [("openssl", "focal")]),
                      ("not-affected", [("openssl1.0", "jammy")])],
    "CVE-2026-1002": [("fixed", [("linux", "jammy"), ("linux-hwe-6.8", "jammy"), ("linux-firmware", "noble")]),
                      ("vulnerable-wip", [("linux", "noble")]),
                      ("ignored", [("linux-aws", "jammy")])],
    "CVE-2026-1003": [("fixed", [("nodejs", "esm-apps/jammy")]), ("deferred", [("nodejs", "jammy")]),
                      ("needs-evaluation", [("nodejs", "noble")]), ("not-in-release", [("nodejs", "esm-infra/focal")])],
    "CVE-2026-1006": [("fixed", [("linux-hwe-6.8", "jammy")]), ("ignored", [("linux-hwe-6.5", "jammy")])],
    "CVE-2026-1007": [("not-affected", [("zlib", "jammy"), ("zlib", "noble")])],
}


class PeriodTest(unittest.TestCase):
    def setUp(self):
        self.env = TempEnv(UBUNTU_RELEASES=["20.04", "22.04", "24.04"]).__enter__()
        osv_import.import_file(self.env.con, make_osv_tar(os.path.join(self.env.dir.name, "o.tar.xz"), RECS, usns=USNS))
        c = self.env.con
        c.execute("INSERT INTO host(id, hostname, release, active) VALUES (1, 'web01', '24.04', 1)")
        c.execute("INSERT INTO host_package VALUES (1, 'libssl3t64', '3.0.13-0ubuntu3')")
        c.commit()

    def tearDown(self):
        self.env.__exit__()

    def with_vex(self):
        vex_import.import_file(self.env.con, make_vex_tar(os.path.join(self.env.dir.name, "v.tar.xz"), VEX))

    def rows(self, minp="medium", status="", have_vex=True, to="2026-09-30"):
        rows, summary = period.build(self.env.con, "2026-07-01", to, minp, status, have_vex=have_vex)
        return {(r["cve_id"], r["package"]): r for r in rows}, summary

    def test_osv_only(self):
        """沒有 VEX：只分 Fixed／Vulnerable"""
        rows, _ = self.rows(have_vex=False)
        r = rows[("CVE-2026-1001", "openssl")]
        self.assertEqual(r["description"], "openssl 的問題")
        self.assertEqual(r["ubuntu20"], "已修正 1.1.1f-1ubuntu2.24+esm1")      # 20.04 以 Pro 為準、不標 Pro
        self.assertEqual(r["ubuntu22"], "已修正 3.0.2-0ubuntu1.20")
        self.assertEqual(rows[("CVE-2026-1003", "nodejs")]["ubuntu22"], "已修正（Pro） 12.22.9~dfsg-1ubuntu3.6+esm1")
        self.assertEqual(rows[("CVE-2026-1003", "nodejs")]["codes"]["ubuntu24"], "none")
        self.assertEqual(rows[("CVE-2026-1002", "linux")]["codes"]["ubuntu24"], "vulnerable")
        self.assertEqual(rows[("CVE-2026-1002", "linux")]["ubuntu22"], "期末後修正（2026-10-05）")

    def test_as_of_period_end(self):
        self.with_vex()
        rows, _ = self.rows()
        self.assertEqual(rows[("CVE-2026-1002", "linux")]["codes"]["ubuntu22"], "fixed-after")           # USN 在 10/5
        self.assertEqual(rows[("CVE-2026-1002", "linux-hwe-6.8")]["codes"]["ubuntu22"], "fixed")          # USN 在 8/10
        self.assertEqual(rows[("CVE-2026-1002", "linux-firmware")]["codes"]["ubuntu24"], "fixed-nodate")  # 沒有 USN
        self.assertEqual(rows[("CVE-2026-1003", "nodejs")]["codes"]["ubuntu22"], "fixed")                 # Pro 的 USN
        rows, summary = self.rows(status="open")
        self.assertIn(("CVE-2026-1002", "linux"), rows)                                                   # 期末未解決包含期末後修正
        self.assertIn("fixed-after", summary["statuses"])
        rows, _ = self.rows(to="2026-10-31")                                                              # 區間拉長就算已修正
        self.assertEqual(rows[("CVE-2026-1002", "linux")]["codes"]["ubuntu22"], "fixed")

    def test_as_of_lookup_both_channels_and_missing_date(self):
        usn = {("CVE-X", "Ubuntu:20.04:LTS", "linux"): ("5.4.0-1", "USN-A", "2026-08-01"),      # ESM 之前，一般套件庫修正
               ("CVE-X", "Ubuntu:Pro:20.04:LTS", "linux"): ("5.4.0-1", "USN-B", "2026-10-01"),
               ("CVE-Y", "Ubuntu:22.04:LTS", "linux"): ("5.15.0-1", "USN-C", "")}               # USN 沒有日期
        self.assertEqual(period.as_of("fixed", True, "5.4.0-1", "CVE-X", "linux", "20.04", usn, "2026-09-30"),
                         ("fixed", ("5.4.0-1", "USN-A", "2026-08-01"), False))                   # 取最早的那個；一般套件庫修正，不需要 Pro
        self.assertEqual(period.as_of("fixed", False, "5.15.0-1", "CVE-Y", "linux", "22.04", usn, "2026-09-30")[0], "fixed-nodate")
        self.assertEqual(period.as_of("vulnerable", False, "", "CVE-Y", "linux", "22.04", usn, "2026-09-30"), ("vulnerable", None, False))
        # 22.04 一般套件庫修正（pro=False）：不能拿較早的 Pro USN 日期
        usn2 = {("CVE-Z", "Ubuntu:22.04:LTS", "openssl"): ("3.0.2-1", "USN-D", "2026-10-02"),
                ("CVE-Z", "Ubuntu:Pro:22.04:LTS", "openssl"): ("3.0.2-1+esm1", "USN-E", "2026-08-01")}
        self.assertEqual(period.as_of("fixed", False, "3.0.2-1", "CVE-Z", "openssl", "22.04", usn2, "2026-09-30")[0], "fixed-after")
        self.assertEqual(period.as_of("fixed", True, "3.0.2-1+esm1", "CVE-Z", "openssl", "22.04", usn2, "2026-09-30"),
                         ("fixed", ("3.0.2-1+esm1", "USN-E", "2026-08-01"), True))                # Pro 的 USN、Pro 的版本
        self.assertEqual(period.as_of("fixed", True, "3.0.2-1+esm1", "CVE-Z", "openssl", "22.04",
                                      {k: v for k, v in usn2.items() if ":Pro:" in k[1]}, "2026-09-30")[2], True)  # 只有 Pro 有修正

    def test_kernel_summary(self):
        self.with_vex()
        _, summary = self.rows()
        k = {(x["release"], x["series"]): x for x in summary["kernel"]}
        h = k[("22.04", "linux-hwe-6.8")]
        self.assertEqual((h["total"], h["fixed_n"], h["upgrade_to"], h["top_cve"], h["top_usn"], h["top_date"]),
                         (2, 2, "6.8.0-91.92~22.04.1", "CVE-2026-1006", "USN-4", "2026-09-20"))     # 取最高版本、最早的 USN
        self.assertIn("升級到 6.8.0-91.92~22.04.1 即可修補這 2 個", h["sentence"])
        self.assertFalse(h["top_pro"])
        # 同一個最高版本有好幾個 CVE：固定取 USN 最早、再取 CVE 編號最小
        tie = [{"cve_id": c, "priority": "high", "package": "linux", "published": "2026-07-01",
                "codes": {"ubuntu20": "none", "ubuntu22": "fixed", "ubuntu24": "none"},
                "fix": {"ubuntu22": ("5.15.0-200.210", u, d)}} for c, u, d in
               (("CVE-2026-9003", "USN-9", "2026-09-10"), ("CVE-2026-9001", "USN-8", "2026-09-12"), ("CVE-2026-9002", "USN-7", "2026-09-10"))]
        t = period.kernel_summary(tie, "2026-07-01", "2026-09-30")[0][0]
        self.assertEqual((t["top_cve"], t["top_usn"], t["top_same"]), ("CVE-2026-9002", "USN-7", 3))
        self.assertIn("同一版本修正的共 3 個", t["sentence"])
        g = k[("22.04", "linux")]
        self.assertEqual((g["fixed_n"], g["open_n"], g["later_n"], g["upgrade_to"]), (0, 1, 1, ""))
        self.assertIn("尚無已發布的修正", g["sentence"])
        self.assertIn("其中 1 個於期末後才修正", g["sentence"])
        self.assertEqual(k[("24.04", "linux")]["open_n"], 1)                                            # 修正中
        self.assertNotIn(("22.04", "linux-aws"), k)                                                      # 雲端變體不列
        self.assertNotIn(("24.04", "linux-firmware"), k)                                                 # 不是 kernel
        _, later = self.rows(to="2026-10-31")
        g2 = {(x["release"], x["series"]): x for x in later["kernel"]}[("22.04", "linux")]
        self.assertEqual((g2["fixed_n"], g2["upgrade_to"]), (1, "5.15.0-160.170"))
        self.assertTrue(any(d["cve"] == "CVE-2026-1002" and d["usn"] == "USN-2" for d in later["kernel_detail"]))

    def test_previous_period(self):
        self.assertEqual(period.previous_period("2026-07-01", "2026-09-30"), ("2026-04-01", "2026-06-30"))
        self.assertEqual(period.previous_period("2026-01-01", "2026-03-31"), ("2025-10-01", "2025-12-31"))
        self.assertEqual(period.previous_period("2026-07-05", "2026-07-14"), ("2026-06-25", "2026-07-04"))
        self.assertIsNone(period.previous_period("0001-01-01", "2026-07-14"))          # 超出日期範圍：不比較

    def test_overview_and_executive(self):
        self.with_vex()
        _, s = self.rows()
        ov = s["overview"]
        # 22.04 critical：linux 期末後修正、linux-hwe-6.8 已修正 → 這個 CVE 算尚未解決；kernel 類
        self.assertEqual(ov[("ubuntu22", "critical", "kernel")], {"applicable": 1, "not_affected": 0, "fixed": 0, "nodate": 0, "open": 1,
                                                                   "fixed_pro": 0, "later": 1})
        self.assertEqual(ov[("ubuntu22", "critical", "other")]["applicable"], 0)        # linux-firmware 只在 24.04
        self.assertEqual(ov[("ubuntu24", "critical", "other")], {"applicable": 1, "not_affected": 0, "fixed": 1, "nodate": 1, "open": 0,
                                                                  "fixed_pro": 0, "later": 0})
        self.assertEqual(ov[("ubuntu20", "medium", "all")]["applicable"], 0)            # nodejs 20.04＝版本無此套件
        e = s["executive"]
        self.assertIn("critical 1 個（其中 kernel 1）", e["headline"])
        m22 = [m for m in e["matrix"] if m["release"] == "22.04"][0]
        self.assertEqual([(c["level"], c["open"], c["applicable"]) for c in m22["cells"]][:2], [("critical", 1, 1), ("high", 0, 2)])   # high：openssl、linux-hwe-6.8 都在期末前修正
        acts = [(a["release"], a["kind"], a["what"]) for a in e["actions"]]
        self.assertIn(("22.04", "kernel", "kernel（linux-hwe-6.8）升級到 6.8.0-91.92~22.04.1"), acts)
        self.assertIn("22.04", [a[0] for a in acts if a[1] == "apt"])
        self.assertEqual(s["prev_range"], ("2026-04-01", "2026-06-30"))
        c = e["cards"]
        self.assertEqual(c["applicable"], c["fixed"] + c["open"] + c["not_affected"])           # 數字卡要加得起來
        self.assertEqual(m22["cells"][0]["delta"], 1)                                     # 上一季沒有資料
        # 期末後才修正：不列入「風險接受／追蹤」，另列「期末後才釋出修正」；數字卡另外顯示
        self.assertIn(("22.04", "later"), [(a["release"], a["kind"]) for a in e["actions"]])
        self.assertEqual(c["later"], 1)
        self.assertIn("其中 1 個期末後已釋出修正", e["headline"])
        track22 = [a["count"] for a in e["actions"] if a["release"] == "22.04" and a["kind"] == "track"]
        self.assertEqual(sum(track22), sum(ov[("ubuntu22", lv, "all")]["open"] - ov[("ubuntu22", lv, "all")]["later"] for lv in s["levels"]))
        # 只選 critical：標籤跟著變
        _, s2 = self.rows(minp="critical")
        self.assertEqual(s2["executive"]["top_label"], "critical")

    def test_pro_only_fix_not_counted_as_apt(self):
        rows = [{"cve_id": "CVE-P", "priority": "high", "package": "openssl", "published": "2026-07-01",
                 "codes": {"ubuntu20": "none", "ubuntu22": "fixed", "ubuntu24": "none"},
                 "fix": {"ubuntu22": ("3.0.2-1+esm1", "USN-E", "2026-08-01")}, "pro": {"ubuntu22": True}}]
        ov = period.overview(rows, ["high"])
        self.assertEqual((ov[("ubuntu22", "high", "other")]["fixed"], ov[("ubuntu22", "high", "other")]["fixed_pro"]), (1, 1))

    def test_as_of_pairs_version_with_usn(self):
        usn = {("CVE-K", "Ubuntu:20.04:LTS", "linux"): ("5.4.0-150", "USN-A", "2026-08-01"),
               ("CVE-K", "Ubuntu:Pro:20.04:LTS", "linux"): ("5.4.0-200", "USN-B", "2026-10-01")}
        self.assertEqual(period.as_of("fixed", True, "5.4.0-200", "CVE-K", "linux", "20.04", usn, "2026-09-30"),
                         ("fixed", ("5.4.0-150", "USN-A", "2026-08-01"), False))     # 選到一般套件庫的 USN：版本配它、不需要 Pro
        # 版本一律配選到的 USN
        usn3 = {("CVE-L", "Ubuntu:22.04:LTS", "curl"): ("7.81.0-1ubuntu1.20", "USN-F", "2026-08-01")}
        self.assertEqual(period.as_of("fixed", False, "7.81.0-1ubuntu1.19", "CVE-L", "curl", "22.04", usn3, "2026-09-30")[1][0],
                         "7.81.0-1ubuntu1.20")
        # 22.04 先有 Pro 修正（期末前）、一般套件庫期末後才修正：期末時的修正是 Pro 的（版本、USN、Pro 一致）
        usn4 = {("CVE-M", "Ubuntu:Pro:22.04:LTS", "curl"): ("7.81.0-1+esm1", "USN-G", "2026-08-01"),
                ("CVE-M", "Ubuntu:22.04:LTS", "curl"): ("7.81.0-1ubuntu1.21", "USN-H", "2026-10-15")}
        self.assertEqual(period.as_of("fixed", True, "7.81.0-1+esm1", "CVE-M", "curl", "22.04", usn4, "2026-09-30"),
                         ("fixed", ("7.81.0-1+esm1", "USN-G", "2026-08-01"), True))
        # 20.04 一般套件庫的 USN 沒有日期：日期不明，但不需要 Pro
        usn5 = {("CVE-N", "Ubuntu:20.04:LTS", "curl"): ("7.68.0-1ubuntu2.30", "USN-I", "")}
        self.assertEqual(period.as_of("fixed", True, "7.68.0-1ubuntu2.30", "CVE-N", "curl", "20.04", usn5, "2026-09-30")[2], False)

    def test_valid_date_bounds(self):
        for bad in ("0001-01-01", "20261231", "2026-W27-1", "2026-7-1", ""):
            with self.assertRaises(ValueError):
                period.valid_date(bad)
        self.assertEqual(period.valid_date("2026-01-01"), "2026-01-01")

    def test_vex_statuses(self):
        self.with_vex()
        rows, _ = self.rows()
        r = rows[("CVE-2026-1001", "openssl")]
        self.assertEqual(r["codes"], {"ubuntu20": "fixed", "ubuntu22": "fixed", "ubuntu24": "fixed"})
        self.assertEqual(r["ubuntu24"], "已修正 3.0.13-0ubuntu3.16")           # 版本用 OSV 的（第一個修正版本）
        self.assertEqual(rows[("CVE-2026-1001", "openssl1.0")]["codes"]["ubuntu22"], "not-affected")
        n = rows[("CVE-2026-1003", "nodejs")]
        self.assertEqual(n["codes"], {"ubuntu20": "not-in-release", "ubuntu22": "fixed", "ubuntu24": "needs-evaluation"})
        self.assertEqual(n["ubuntu22"], "已修正（Pro） 12.22.9~dfsg-1ubuntu3.6+esm1")
        self.assertEqual(rows[("CVE-2026-1002", "linux")]["ubuntu24"], "修正中")
        self.assertEqual(rows[("CVE-2026-1001", "openssl1.0")]["ubuntu24"], "—")
        # 20.04 的 focal 已不支援（Ignored），但 esm-infra 有修正 → 以 Pro 為準
        self.assertEqual(rows[("CVE-2026-1001", "openssl")]["ubuntu20"], "已修正 1.1.1f-1ubuntu2.24+esm1")

    def test_cve_without_packages_listed(self):
        self.with_vex()
        rows, summary = self.rows()
        z = rows[("CVE-2026-1007", "zlib")]
        self.assertEqual(z["codes"], {"ubuntu20": "none", "ubuntu22": "not-affected", "ubuntu24": "not-affected"})
        rows, _ = self.rows(have_vex=False)          # 沒有 VEX 也要列出 CVE（套件欄空白）
        self.assertIn(("CVE-2026-1007", ""), rows)

    def test_kernel_filter(self):
        self.with_vex()
        rows, _ = self.rows()
        pk = sorted(p for c, p in rows if c == "CVE-2026-1002")
        self.assertEqual(pk, ["linux", "linux-firmware", "linux-hwe-6.8"])   # 雲端變體不列；22.04 最新的 HWE
        self.assertNotIn(("CVE-2026-1006", "linux-hwe-6.5"), rows)            # 舊 HWE 不列
        d = rows[("CVE-2026-1006", "linux-hwe-6.8")]["description"]
        self.assertTrue(d.startswith("x" * period.DESC_MAX) and "ubuntu.com/security/CVE-2026-1006" in d)

    def test_summary_worst_status(self):
        self.with_vex()
        rows, summary = self.rows()
        self.assertEqual({c for c, _ in rows}, {"CVE-2026-1001", "CVE-2026-1002", "CVE-2026-1003", "CVE-2026-1006", "CVE-2026-1007"})
        self.assertEqual(summary["total"], 5)
        crit = summary["by_level"]["critical"]
        self.assertEqual((crit["cves"], crit["kernel"]), (1, 1))
        self.assertEqual(crit["ubuntu24"]["vulnerable-wip"], 1)               # linux（wip）比 linux-firmware（fixed）優先
        self.assertEqual(crit["ubuntu20"]["none"], 1)
        med = summary["by_level"]["medium"]
        self.assertEqual(med["ubuntu24"]["needs-evaluation"], 1)
        self.assertEqual(med["ubuntu24"]["not-affected"], 1)
        self.assertNotIn("unknown", summary["statuses"])
        rows, _ = self.rows(minp="low")
        self.assertIn(("CVE-2026-1005", "curl"), rows)

    def test_status_filter_and_sort(self):
        self.with_vex()
        rows, summary = self.rows(status="open")
        self.assertEqual(sorted(rows), [("CVE-2026-1002", "linux"), ("CVE-2026-1003", "nodejs")])
        self.assertEqual(summary["total"], 5)                                 # 篩選不影響統計
        rows, _ = self.rows(status="needs-evaluation")
        self.assertEqual(list(rows), [("CVE-2026-1003", "nodejs")])
        lst, _ = period.build(self.env.con, "2026-07-01", "2026-09-30", "medium")
        period.sort_rows(lst, "ubuntu24")
        self.assertEqual(lst[0]["codes"]["ubuntu24"], "needs-evaluation")    # 依狀態順序，最需要注意的在前

    def test_xlsx(self):
        from openpyxl import load_workbook
        self.with_vex()
        rows, summary = period.build(self.env.con, "2026-07-01", "2026-09-30", "medium")
        wb = load_workbook(io.BytesIO(period.to_xlsx(rows, summary, "2026-07-01", "2026-09-30", "medium", vex_date="2026-10-21")))
        self.assertEqual(wb.sheetnames, ["主管摘要", "統計", "狀態明細", "Kernel 修補版本", "Kernel 明細", "狀態說明", "明細"])
        self.assertIn("公開的弱點", wb["主管摘要"]["A2"].value)
        kn = [[c.value for c in r] for r in wb["Kernel 修補版本"].iter_rows(min_row=3)]
        self.assertTrue(any(r[1] == "linux-hwe-6.8" and r[5] == "6.8.0-91.92~22.04.1" for r in kn))
        head = [c.value for c in wb["明細"][1]]
        self.assertEqual(head, ["CVE", "等級", "公開日期", "套件", "Ubuntu 20.04", "Ubuntu 22.04", "Ubuntu 24.04", "說明"])
        self.assertEqual(wb["明細"].max_row, len(rows) + 1)
        self.assertTrue(any(c.value == "待評估" for c in wb["狀態明細"][9]))
        self.assertEqual([c.value for c in wb["狀態說明"][2]], ["待評估", "官方還沒判斷是否受影響", "Needs evaluation"])


if __name__ == "__main__":
    unittest.main()
