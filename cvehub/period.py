"""輸出報表（季報表）：依 CVE 公開日期（OSV 的 published），列出 Ubuntu 官方對每個 CVE × 套件 × 版本的處理狀態

跟 cvehub 管理的主機無關（主機受影響情形見 impact.py）。

- 一列＝一個 CVE × 一個 source 套件；沒有任何相關套件的 CVE 也列一列（套件欄空白），CVE 數才會完整
- 狀態來自 Canonical VEX（vex_import.STATUSES，畫面與 Excel 用繁體中文）；還沒匯入 VEX 時退回 OSV，只分已修正／受影響
- ubuntu20 以 Ubuntu Pro（esm）為準（20.04 標準支援已結束）；ubuntu22／ubuntu24 以一般套件庫為準，只有 Pro 修正時標（Pro）
- 修正版本用 OSV 的（第一個修正版本）
- kernel 只留 linux 與 linux-hwe-*，其他 kernel 變體（aws、azure、gcp、oem、lowlatency…）不列；
  HWE 只留各版本最新的系列，以及主機上實際有裝的系列（停止支援的舊 HWE 永遠未修，列出來只會誤導）
- 統計：一個 CVE 在某版本有多個套件時，取最需要注意的狀態（vex_import.STATUSES 的順序）
"""
import datetime
import re
import functools
import io

from .config import PRIO
from .kernel import is_kernel, kernel_variant
from .sync import KERNEL_ABI, _kernel_source_by_series
from .vercmp import vercmp
from .vex_import import LABEL, LABEL_EN, RANK, STATUSES, UNRESOLVED, status_for

# 報表期末時尚未修正（含期末後才修正的）
OPEN_AT_END = tuple(UNRESOLVED) + ("fixed-after",)

RELEASES = [("ubuntu20", "20.04"), ("ubuntu22", "22.04"), ("ubuntu24", "24.04")]
COLUMNS = ["cve_id", "priority", "published", "package", "ubuntu20", "ubuntu22", "ubuntu24", "description"]
HEADERS = {"cve_id": "CVE", "priority": "等級", "published": "公開日期", "package": "套件",
           "ubuntu20": "Ubuntu 20.04", "ubuntu22": "Ubuntu 22.04", "ubuntu24": "Ubuntu 24.04", "description": "說明"}
LEVELS = ["critical", "high", "medium", "low", "negligible"]
DESC_MAX = 1000                  # kernel CVE 的說明常是整段 commit message
STATUS_FILTERS = [("", "全部"), ("open", "期末未解決（待評估／受影響／修正中／延後修正／不修正／期末後修正）")] + \
                 [(c, zh) for c, _, zh, _ in STATUSES if c != "unknown"]
FILLS = {"needs-evaluation": "FFB3AD", "vulnerable": "FFD8A8", "vulnerable-wip": "FFE8A3", "deferred": "FFF1B8",
         "ignored": "E1E4E8", "fixed-after": "FFD8A8", "unknown": "E1E4E8", "fixed-nodate": "E6F4EA", "fixed": "D3F9D8",
         "not-affected": "F6F8FA", "not-in-release": "F6F8FA", "none": "FFFFFF"}
PRIO_FILLS = {"critical": "CF222E", "high": "FB8500", "medium": "FFF1B8", "low": "DDF4FF", "negligible": "EEEEEE"}


def _hwe_series(src):
    try:
        return tuple(int(x) for x in src[len("linux-hwe-"):].split("."))
    except ValueError:
        return ()


def _ecosystems(rel):
    """(主要, 備用)；20.04 只看 Pro"""
    if rel == "20.04":
        return "Ubuntu:Pro:20.04:LTS", None
    return "Ubuntu:%s:LTS" % rel, "Ubuntu:Pro:%s:LTS" % rel


def latest_hwe(con):
    """各版本 OSV 資料裡最新的 HWE 系列，例：{"22.04": "linux-hwe-6.8"}"""
    out = {}
    for col, rel in RELEASES:
        main, _ = _ecosystems(rel)
        srcs = [r[0] for r in con.execute("SELECT DISTINCT source FROM affected WHERE ecosystem=? AND source LIKE 'linux-hwe-%'", (main,))]
        if srcs:
            out[rel] = max(srcs, key=_hwe_series)
    return out


def _desc(cve, text):
    text = (text or "").strip()
    if len(text) <= DESC_MAX:
        return text
    return text[:DESC_MAX].rstrip() + "…（全文：https://ubuntu.com/security/%s）" % cve


def host_sources(con):
    """{host_id: (release, {source})}：用和比對相同的 binary→source 對應"""
    out = {}
    for h in con.execute("SELECT id, release FROM host WHERE active=1"):
        rel = h["release"]
        ecos = ["Ubuntu:%s:LTS" % rel, "Ubuntu:Pro:%s:LTS" % rel]
        names = [r["name"] for r in con.execute("SELECT name FROM host_package WHERE host_id=?", (h["id"],))]
        found, srcs = set(), set()
        for i in range(0, len(names), 500):
            chunk = names[i:i + 500]
            q = "SELECT binary, source FROM binmap WHERE ecosystem IN (%s) AND binary IN (%s)" % (
                ",".join("?" * len(ecos)), ",".join("?" * len(chunk)))
            for r in con.execute(q, ecos + chunk):
                found.add(r["binary"]); srcs.add(r["source"])
        for n in names:                  # 新版 kernel 還沒出現在 OSV 的 binary 清單時，用同系列推測
            if n not in found and KERNEL_ABI.match(n):
                for e in ecos:
                    s = _kernel_source_by_series(con, e, n)
                    if s:
                        srcs.add(s); break
        out[h["id"]] = (rel, srcs)
    return out


def cell_text(code, pro, fixed, release, fix_date=""):
    if code == "fixed-after":
        return "期末後修正（%s）" % fix_date
    t = "—" if code == "none" else LABEL.get(code, code)
    if code in ("fixed", "fixed-nodate"):
        if pro and release != "20.04":
            t += "（Pro）"
        if fixed:
            t += " " + fixed
    return t


def usn_fixes(con, cves):
    """{(cve, ecosystem, source): (修正版本, USN, 發布日)}：每個 CVE 第一次被修正的 USN"""
    out = {}
    for i in range(0, len(cves), 500):
        chunk = cves[i:i + 500]
        for r in con.execute("SELECT * FROM usn_fix WHERE cve IN (%s)" % ",".join("?" * len(chunk)), chunk):
            out[(r["cve"], r["ecosystem"], r["source"])] = (r["fixed"], r["usn"], r["published"])
    return out


def as_of(code, pro, fx, cve, src, rel, usn, date_to):
    """「已修正」改以報表期末判斷：USN 在期末之後＝期末後修正；沒有 USN（或 USN 沒有日期）＝日期不明。
    回傳 (代碼, 修正資訊 (版本, USN, 發布日), 這個修正是否來自 Pro)

    - 一般套件庫修正（pro=False）只看一般套件庫的 USN；靠 Pro 修正（pro=True，20.04 有 Pro 狀態時都是）兩邊都看、
      取最早有日期的那個：例如 20.04 的 CVE 可能在 ESM 之前就由一般套件庫的 USN 修正
    - 版本、USN、日期、是否 Pro 全部來自選到的同一個 USN（描述「期末時是哪個修正」，評鑑單引用時不會對不上）
    - 沒有日期的：一般套件庫有 USN（沒日期）就不算 Pro，否則沿用傳入的 pro"""
    if code != "fixed":
        return code, None, pro
    base_eco, pro_eco = "Ubuntu:%s:LTS" % rel, "Ubuntu:Pro:%s:LTS" % rel
    cands = []
    for eco, is_pro in [(base_eco, False)] + ([(pro_eco, True)] if pro else []):
        u = usn.get((cve, eco, src))
        if u and u[2]:
            cands.append((u, is_pro))
    if not cands:
        return "fixed-nodate", (fx, "", ""), pro and not usn.get((cve, base_eco, src))
    u, from_pro = min(cands, key=lambda c: c[0][2])
    return ("fixed-after" if u[2] > date_to else "fixed"), (u[0] or fx, u[1], u[2]), from_pro


def osv_fixed(con, cves, releases):
    """{(cve, source, release): {"base": 修正版本或 "", "pro": ...}}；有列＝OSV 認定受影響"""
    ecos = {}
    for rel in releases:
        ecos["Ubuntu:%s:LTS" % rel] = (rel, "base")
        ecos["Ubuntu:Pro:%s:LTS" % rel] = (rel, "pro")
    out = {}
    for i in range(0, len(cves), 500):
        chunk = cves[i:i + 500]
        # +ecosystem：讓 SQLite 用 cve 索引（不加會改走 ecosystem 索引，慢約 90 倍）
        q = "SELECT ecosystem, source, cve, fixed FROM affected WHERE cve IN (%s) AND +ecosystem IN (%s)" % (
            ",".join("?" * len(chunk)), ",".join("?" * len(ecos)))
        for r in con.execute(q, chunk + list(ecos)):
            rel, ch = ecos[r["ecosystem"]]
            out.setdefault((r["cve"], r["source"], rel), {})[ch] = r["fixed"] or ""
    return out


def _osv_only_status(fx, release):
    """沒有 VEX 時由 OSV 推：有修正版本＝fixed，否則 vulnerable"""
    if not fx:
        return "none", False, ""
    main, other = ("pro", "base") if release == "20.04" else ("base", "pro")
    if fx.get(main):
        return "fixed", main == "pro", fx[main]
    if fx.get(other):
        return "fixed", other == "pro", fx[other]
    return "vulnerable", False, ""


def resolve(vex_row, fx, release, have_vex):
    """一格的 (狀態代碼, 是否靠 Pro, 修正版本)"""
    if have_vex and vex_row is not None:
        return status_for(vex_row, release, fx)
    return _osv_only_status(fx, release)


def build(con, date_from, date_to, min_priority="medium", status_filter="", have_vex=True, compare=True):
    """回傳 (rows, summary)。date_from／date_to 是 YYYY-MM-DD（含頭尾）；compare＝主管摘要要不要和上一期比較"""
    maxp = PRIO.get(min_priority, 2)
    levels = [lv for lv in LEVELS if PRIO[lv] <= maxp]
    rels = [rel for _, rel in RELEASES]
    vulns = {r["cve"]: r for r in con.execute(
        "SELECT cve, priority, details, published FROM vuln WHERE substr(published,1,10) BETWEEN ? AND ? "
        "AND priority IN (%s)" % ",".join("?" * len(levels)), [date_from, date_to] + levels)}
    cves = list(vulns)
    fixed = osv_fixed(con, cves, rels)
    vex = {}
    if have_vex:
        for i in range(0, len(cves), 500):
            chunk = cves[i:i + 500]
            q = "SELECT * FROM vex_status WHERE cve IN (%s) AND release IN (%s)" % (
                ",".join("?" * len(chunk)), ",".join("?" * len(rels)))
            for r in con.execute(q, chunk + rels):
                vex[(r["cve"], r["source"], r["release"])] = r
    usn = usn_fixes(con, cves)
    keys = {(c, s) for (c, s, _) in fixed} | {(c, s) for (c, s, _) in vex}
    keep_hwe = set(latest_hwe(con).values()) | {s for _, srcs in host_sources(con).values() for s in srcs if s.startswith("linux-hwe-")}
    keys = {(c, s) for (c, s) in keys if not kernel_variant(s) and not (s.startswith("linux-hwe-") and s not in keep_hwe)}
    with_rows = {c for c, _ in keys}
    keys |= {(c, "") for c in cves if c not in with_rows}      # 沒有相關套件的 CVE 也列

    rows = []
    for cve, src in keys:
        v = vulns[cve]
        row = {"cve_id": cve, "priority": v["priority"], "published": (v["published"] or "")[:10],
               "description": _desc(cve, v["details"]), "package": src, "codes": {}, "fix": {}, "pro": {}}
        for col, rel in RELEASES:
            code, pro, fx = resolve(vex.get((cve, src, rel)), fixed.get((cve, src, rel)), rel, have_vex) if src else ("none", False, "")
            code, fix, pro = as_of(code, pro, fx, cve, src, rel, usn, date_to)
            row["codes"][col], row["fix"][col], row["pro"][col] = code, fix, pro
            row[col] = cell_text(code, pro, fix[0] if fix else fx, rel, fix[2] if fix else "")
        rows.append(row)

    summary = _summary(rows, levels)
    summary["kernel"], summary["kernel_detail"] = kernel_summary(rows, date_from, date_to)
    summary["overview"] = overview(rows, levels)
    prev_ov = None
    if compare:
        prev = previous_period(date_from, date_to)
        if prev:
            summary["prev_range"] = prev
            prev_ov = _prev_overview(con, prev[0], prev[1], min_priority, have_vex)
    summary["executive"] = executive(summary, date_from, date_to, prev_ov)
    if status_filter == "open":
        rows = [r for r in rows if any(c in OPEN_AT_END for c in r["codes"].values())]
    elif status_filter:
        rows = [r for r in rows if status_filter in r["codes"].values()]
    rows.sort(key=lambda x: (PRIO.get(x["priority"], 5), x["cve_id"], x["package"]))
    return rows, summary


_PREV_CACHE = {}


def _prev_overview(con, pf, pt, min_priority, have_vex):
    """上一期的 overview（只用來算 ▲▼）。換頁、排序、匯出都會重算報表，上一期結果依資料版本快取，不必每次重建"""
    # 鍵：資料庫檔案＋使用中的 OSV／VEX＋最後一次結束的同步（成功或失敗都算：失敗的同步也可能已更新部分主機）
    #     ＋使用中的主機（數量、id 合計、最後同步時間；主機的 HWE 會影響 kernel 列）
    db = next((r[2] for r in con.execute("PRAGMA database_list") if r[1] == "main"), "")
    key = (db, pf, pt, min_priority, have_vex) + tuple(con.execute(
        "SELECT (SELECT max(id) FROM osv_import WHERE active=1), (SELECT max(id) FROM vex_import WHERE active=1), "
        "(SELECT max(id) FROM sync_run WHERE finished_at IS NOT NULL), "
        "(SELECT count(*) FROM host WHERE active=1), (SELECT total(id) FROM host WHERE active=1), "
        "(SELECT max(last_seen) FROM host WHERE active=1)").fetchone())
    ov = _PREV_CACHE.get(key)
    if ov is None:
        ov = build(con, pf, pt, min_priority, "", have_vex, compare=False)[1]["overview"]
        if len(_PREV_CACHE) >= 16:
            _PREV_CACHE.clear()
        _PREV_CACHE[key] = ov
    return ov


def _summary(rows, levels):
    """每個等級：CVE 數、kernel CVE 數；每個版本：以 CVE 計的狀態分布（多個套件取最需要注意的）"""
    worst = {}                                  # (cve, col) -> code
    prio, kernel = {}, set()
    for r in rows:
        prio[r["cve_id"]] = r["priority"]
        if is_kernel(r["package"]):
            kernel.add(r["cve_id"])
        for col, code in r["codes"].items():
            k = (r["cve_id"], col)
            if k not in worst or RANK[code] < RANK[worst[k]]:
                worst[k] = code
    by_level = {}
    for lv in levels:
        mine = [c for c, p in prio.items() if p == lv]
        s = {"cves": len(mine), "kernel": len([c for c in mine if c in kernel])}
        for col, _ in RELEASES:
            cnt = {code: 0 for code, _, _, _ in STATUSES}
            for c in mine:
                cnt[worst.get((c, col), "none")] += 1
            s[col] = cnt
        by_level[lv] = s
    optional = ("unknown", "fixed-after", "fixed-nodate")        # 有數字才顯示這一欄
    used = [code for code, _, _, _ in STATUSES
            if code not in optional or any(by_level[lv][col][code] for lv in levels for col, _ in RELEASES)]
    return {"levels": levels, "by_level": by_level, "statuses": used, "total": len(prio), "total_kernel": len(kernel)}


_SERIES_ORDER = lambda s: (0, ()) if s == "linux" else (1, _hwe_series(s))


def kernel_summary(rows, date_from, date_to):
    """每個 Ubuntu 版本 × kernel 系列：期末前已修正的 kernel CVE 要升級到哪個版本（取最高的修正版本）

    給風險評鑑單用：本期（CVE 公開日期在區間內）的 kernel 弱點有幾個、截至期末升級到哪一版可以修補、
    最高的修正版本來自哪個 CVE／USN；期末仍未修正與修正日期不明的另外列出。回傳 (摘要列, 明細列)"""
    groups, detail = {}, []
    for r in rows:
        if not is_kernel(r["package"]):
            continue
        for col, rel in RELEASES:
            code = r["codes"][col]
            if code in ("none", "not-affected", "not-in-release"):
                continue
            groups.setdefault((rel, r["package"]), []).append((r, col, code))
            fix = r["fix"].get(col) or ("", "", "")
            detail.append({"release": rel, "series": r["package"], "cve": r["cve_id"], "priority": r["priority"],
                           "published": r["published"], "status": LABEL.get(code, code), "code": code,
                           "fixed": fix[0] or "", "usn": fix[1] or "", "fix_date": fix[2] or ""})
    out = []
    for (rel, src) in sorted(groups, key=lambda k: (k[0], _SERIES_ORDER(k[1]))):
        items = groups[(rel, src)]
        prio = {}
        for r, col, code in items:
            prio[r["cve_id"]] = r["priority"]
        fixed = [(r["fix"][col][0], r["cve_id"], r["fix"][col][1], r["fix"][col][2], bool(r.get("pro", {}).get(col)))
                 for r, col, code in items if code == "fixed"]
        later = sorted({(r["fix"][col][2], r["cve_id"]) for r, col, code in items if code == "fixed-after"})
        still = sorted({r["cve_id"] for r, col, code in items if code in UNRESOLVED})
        nodate = sorted({r["cve_id"] for r, col, code in items if code == "fixed-nodate"})
        top, same = None, 0
        if fixed:
            best = max(fixed, key=functools.cmp_to_key(lambda a, b: vercmp(a[0], b[0])))[0]
            tied = [f for f in fixed if vercmp(f[0], best) == 0]
            top = min(tied, key=lambda f: (f[3], f[1]))           # 同版本多個：取 USN 最早、再取 CVE 編號最小（固定、可重現）
            same = len({f[1] for f in tied})
        by_lv = [(lv, sum(1 for p in prio.values() if p == lv)) for lv in LEVELS]
        by_lv = [(lv, n) for lv, n in by_lv if n]
        open_n = len(later) + len(still)
        s = "Ubuntu %s（%s）：%s～%s 公開的 kernel 弱點共 %d 個（%s）" % (
            rel, src, date_from, date_to, len(prio), "、".join("%s %d" % kv for kv in by_lv))
        if top:
            pro_only = len({f[1] for f in fixed if f[4]})
            s += "，截至 %s 已修正 %d 個，升級到 %s%s 即可修補這 %d 個（取最高者 %s，%s，%s 發布%s）" % (
                date_to, len({f[1] for f in fixed}), top[0],
                "（需要 Ubuntu Pro）" if top[4] else "（其中 %d 個的修正來自 Ubuntu Pro）" % pro_only if pro_only else "",
                len({f[1] for f in fixed}), top[1], top[2], top[3],
                "；同一版本修正的共 %d 個" % same if same > 1 else "")
        else:
            s += "，截至 %s 尚無已發布的修正" % date_to
        if open_n:
            s += "；期末仍未修正 %d 個" % open_n
            if later:
                s += "（其中 %d 個於期末後才修正）" % len(later)
        if nodate:
            s += "；另有 %d 個有修正版本但無法確認修正日期" % len(nodate)
        out.append({"release": rel, "series": src, "total": len(prio), "by_level": by_lv,
                    "fixed_n": len({f[1] for f in fixed}), "upgrade_to": top[0] if top else "",
                    "top_cve": top[1] if top else "", "top_usn": top[2] if top else "", "top_date": top[3] if top else "",
                    "top_same": same, "top_pro": bool(top and top[4]),
                    "open_n": open_n, "later_n": len(later), "nodate_n": len(nodate), "sentence": s + "。"})
    detail.sort(key=lambda d: (d["release"], _SERIES_ORDER(d["series"]), RANK.get(d["code"], 99), PRIO.get(d["priority"], 5), d["cve"]))
    return out, detail


KINDS = [("kernel", "kernel"), ("other", "其他套件")]
UNRELATED = ("none", "not-in-release")


def _bucket(code):
    """主管摘要用的四類：無關／不受影響／已修正（期末前，含日期不明）／尚未解決（含期末後修正）"""
    if code in UNRELATED:
        return "unrelated"
    if code == "not-affected":
        return "not_affected"
    if code in ("fixed", "fixed-nodate"):
        return "fixed"
    return "open"


def overview(rows, levels):
    """{(col, level, kind): {"applicable","not_affected","fixed","nodate","open"}}；kind＝kernel／other／all

    以 CVE 計：同一個 CVE 在某版本有多個套件時，取最需要注意的狀態（同統計表）"""
    worst = {}                      # (col, kind, cve) -> (code, 是否靠 Pro 修正)
    prio = {}
    for r in rows:
        prio[r["cve_id"]] = r["priority"]
        kind = "kernel" if is_kernel(r["package"]) else "other"
        for col, code in r["codes"].items():
            pro = bool(r.get("pro", {}).get(col))
            for k in (kind, "all"):
                key = (col, k, r["cve_id"])
                if key not in worst or RANK[code] < RANK[worst[key][0]] or (code == worst[key][0] and pro and not worst[key][1]):
                    worst[key] = (code, pro)
    out = {}
    for col, _ in RELEASES:
        for lv in levels:
            for kind in ("kernel", "other", "all"):
                out[(col, lv, kind)] = {"applicable": 0, "not_affected": 0, "fixed": 0, "nodate": 0, "open": 0,
                                        "fixed_pro": 0, "later": 0}
    for (col, kind, cve), (code, pro) in worst.items():
        lv = prio[cve]
        if (col, lv, kind) not in out:
            continue
        b = _bucket(code)
        if b == "unrelated":
            continue
        c = out[(col, lv, kind)]
        c["applicable"] += 1
        c[b] += 1
        if code == "fixed-nodate":
            c["nodate"] += 1
        if b == "fixed" and pro:
            c["fixed_pro"] += 1        # 要 Ubuntu Pro（ESM）才拿得到修正，apt upgrade 不會裝
        if code == "fixed-after":
            c["later"] += 1            # 期末時尚未修正，但現在已有修正
    return out


def previous_period(date_from, date_to):
    """比較用的上一期：區間剛好是整季就取上一季，否則取前一段等長的區間；算出來超出日期範圍時回傳 None（不比較）"""
    f, t = datetime.date.fromisoformat(date_from), datetime.date.fromisoformat(date_to)
    q_start = f.month in (1, 4, 7, 10) and f.day == 1
    q_end = (t + datetime.timedelta(days=1)).day == 1 and t.month == f.month + 2 and t.year == f.year
    try:
        pt = f - datetime.timedelta(days=1)
        pf = datetime.date(pt.year, pt.month - 2, 1) if (q_start and q_end) else pt - (t - f)
    except OverflowError:
        return None
    return pf.isoformat(), pt.isoformat()


def executive(summary, date_from, date_to, prev_overview=None):
    """主管摘要：一句話結論、數字卡（選定等級中的 critical＋high）、版本 × 等級矩陣（尚未解決／適用，含與上一期比較）、行動表"""
    ov, levels = summary["overview"], summary["levels"]
    top = [lv for lv in ("critical", "high") if lv in levels]
    matrix = []
    for col, rel in RELEASES:
        cells = []
        for lv in levels:
            c, k, o = ov[(col, lv, "all")], ov[(col, lv, "kernel")], ov[(col, lv, "other")]
            delta = None
            if prev_overview is not None:
                delta = c["open"] - prev_overview[(col, lv, "all")]["open"]
            cells.append({"level": lv, "open": c["open"], "applicable": c["applicable"], "fixed": c["fixed"],
                          "kernel_open": k["open"], "other_open": o["open"], "delta": delta})
        matrix.append({"col": col, "release": rel, "cells": cells})
    def total(field, kind="all"):
        return sum(ov[(col, lv, kind)][field] for col, _ in RELEASES for lv in top)
    cards = {"applicable": total("applicable"), "not_affected": total("not_affected"), "fixed": total("fixed"),
             "open": total("open"), "later": total("later")}
    # 行動表：每個版本，kernel 升級（每個系列）＋一般套件可更新數（一般套件庫／需 Pro 分開）＋期末後才修正＋官方未修正（選定的所有等級）
    actions = []
    for col, rel in RELEASES:
        for k in summary["kernel"]:
            if k["release"] == rel and k["upgrade_to"]:
                actions.append({"release": rel, "what": "kernel（%s）升級到 %s%s" % (k["series"], k["upgrade_to"],
                                                                          "（需要 Ubuntu Pro）" if k["top_pro"] else ""),
                                "count": k["fixed_n"], "kind": "kernel"})
        other_fixed = sum(ov[(col, lv, "other")]["fixed"] for lv in levels)
        other_pro = sum(ov[(col, lv, "other")]["fixed_pro"] for lv in levels)
        if other_fixed - other_pro:
            actions.append({"release": rel, "what": "一般套件更新（apt upgrade）", "count": other_fixed - other_pro, "kind": "apt"})
        if other_pro:
            actions.append({"release": rel, "what": "需要 Ubuntu Pro（ESM）的套件更新（主機要先啟用 Pro）", "count": other_pro, "kind": "pro"})
        later = sum(ov[(col, lv, "all")]["later"] for lv in levels)
        if later:
            actions.append({"release": rel, "what": "期末後才釋出修正（現在已可更新，下一期報表會列為已修正）", "count": later,
                            "kind": "later"})
        still = sum(ov[(col, lv, "all")]["open"] - ov[(col, lv, "all")]["later"] for lv in levels)
        if still:
            actions.append({"release": rel, "what": "官方尚未修正：風險接受／持續追蹤", "count": still, "kind": "track"})
    by = summary["by_level"]
    top_label = "＋".join(top) if top else "critical＋high"
    parts = ["%s %d 個（其中 kernel %d）" % (lv, by[lv]["cves"], by[lv]["kernel"]) for lv in top]
    open_txt = "、".join("%s %d" % (m["release"], sum(c["open"] for c in m["cells"] if c["level"] in top)) for m in matrix)
    headline = "%s～%s 公開的弱點：%s。截至 %s，%s 尚未解決：%s%s。官方已釋出修正的，可依「要做的事」升級 kernel 與更新套件處理；其餘官方尚未修正，列入追蹤。" % (
        date_from, date_to, "、".join(parts) if parts else "（無 critical／high）", date_to, top_label, open_txt,
        "（其中 %d 個期末後已釋出修正）" % cards["later"] if cards["later"] else "")
    return {"headline": headline, "cards": cards, "matrix": matrix, "actions": actions, "top_levels": top, "top_label": top_label}


def sort_rows(rows, key, desc=False):
    """依欄位排序（穩定排序；等級依嚴重度、版本欄依狀態順序、其他依字串）"""
    if key == "priority":
        k = lambda x: PRIO.get(x["priority"], 5)
    elif key in dict(RELEASES):
        k = lambda x: (RANK.get(x["codes"][key], 99), x[key])
    else:
        k = lambda x: str(x[key] or "").lower()
    rows.sort(key=k, reverse=desc)


def to_xlsx(rows, summary, date_from, date_to, min_priority, osv_date=None, vex_date=None, status_filter=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ex = wb.active
    ex.title = "主管摘要"
    _xlsx_executive(ex, summary, date_from, date_to, osv_date, vex_date)
    ov = wb.create_sheet("統計")
    _xlsx_overview(ov, summary)
    ws = wb.create_sheet("狀態明細")
    info = [("報表區間（CVE 公開日期）", "%s ～ %s" % (date_from, date_to)),
            ("等級", "、".join(summary["levels"])),
            ("狀態篩選（只影響明細）", dict(STATUS_FILTERS).get(status_filter, status_filter) if status_filter else "全部"),
            ("OSV 資料日期", osv_date or ""),
            ("VEX 資料日期（官方狀態）", vex_date or "未匯入：狀態只分「已修正／受影響」"),
            ("產生時間", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("說明", "跟主機無關，是 Ubuntu 官方對每個 CVE 的處理狀態。「已修正」以報表期末日判斷（修正的 USN 在期末前發布）；"
             "期末之後才修正的列為「期末後修正」，沒有 USN 的列為「已修正（日期不明）」。CVE 數以 CVE 編號計；"
             "一個 CVE 在某版本有多個套件時，以最需要注意的狀態計（順序見「狀態說明」）。"
             "20.04 以 Ubuntu Pro 為準；22.04／24.04 只有 Pro 修正時標（Pro）。「—」＝官方資料沒有這個套件／版本的紀錄。"
             "kernel 只列 linux、各版本最新的 HWE 與主機上有裝的 HWE。")]
    for k, v in info:
        ws.append([k, v])
    ws.append([])
    head = ["等級", "版本", "本期 CVE 數（該等級全部）"] + [LABEL[c] for c in summary["statuses"]]
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for i, code in enumerate(summary["statuses"]):
        ws.cell(ws.max_row, 4 + i).fill = PatternFill("solid", fgColor=FILLS[code])
    for col, rel in RELEASES:                      # 版本優先，再分等級
        for lv in summary["levels"]:
            s = summary["by_level"][lv]
            ws.append([lv, "Ubuntu " + rel, s["cves"]] + [s[col][c] for c in summary["statuses"]])
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=PRIO_FILLS[lv])
    ws.append(["合計 CVE", "", summary["total"]])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["其中 kernel", "", summary["total_kernel"]])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14
    for i in range(3, len(head) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16

    kn = wb.create_sheet("Kernel 修補版本")
    kn.append(["Kernel 修補版本（截至 %s）：每個 Ubuntu 版本 × kernel 系列，期末前已發布的修正要升級到哪個版本（取最高的修正版本）" % date_to])
    kn.cell(1, 1).font = Font(bold=True)
    kn.append(["Ubuntu", "kernel 系列", "kernel 弱點數", "等級", "期末前已修正", "升級到", "最高者 CVE", "USN", "USN 發布日",
               "期末仍未修正", "其中期末後才修正", "修正日期不明", "評鑑單文字"])
    for c in kn[2]:
        c.font = Font(bold=True)
    for k in summary.get("kernel", []):
        kn.append([k["release"], k["series"], k["total"], "、".join("%s %d" % kv for kv in k["by_level"]), k["fixed_n"],
                   k["upgrade_to"] + ("（需要 Ubuntu Pro）" if k["top_pro"] else ""), k["top_cve"], k["top_usn"], k["top_date"], k["open_n"], k["later_n"], k["nodate_n"], k["sentence"]])
        kn.cell(kn.max_row, 13).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKLM", (9, 16, 12, 26, 12, 24, 17, 13, 12, 12, 14, 12, 90)):
        kn.column_dimensions[col].width = w
    kd = wb.create_sheet("Kernel 明細")
    kd.append(["Ubuntu", "kernel 系列", "CVE", "等級", "公開日期", "期末狀態", "修正版本", "USN", "USN 發布日"])
    for c in kd[1]:
        c.font = Font(bold=True)
    for d in summary.get("kernel_detail", []):
        kd.append([d["release"], d["series"], d["cve"], d["priority"], d["published"], d["status"], d["fixed"], d["usn"], d["fix_date"]])
    for col, w in zip("ABCDEFGHI", (9, 16, 17, 10, 11, 18, 26, 13, 12)):
        kd.column_dimensions[col].width = w
    kd.freeze_panes = "A2"
    kd.auto_filter.ref = kd.dimensions

    lg = wb.create_sheet("狀態說明")
    lg.append(["狀態", "說明", "Ubuntu 官方名稱"])
    for c in lg[1]:
        c.font = Font(bold=True)
    for code, en, zh, desc in STATUSES:
        lg.append([zh, desc, en])
        lg.cell(lg.max_row, 1).fill = PatternFill("solid", fgColor=FILLS[code])
    lg.column_dimensions["A"].width = 16
    lg.column_dimensions["B"].width = 60
    lg.column_dimensions["C"].width = 30

    d = wb.create_sheet("明細")
    d.append([HEADERS.get(c, c) for c in COLUMNS])
    for c in d[1]:
        c.font = Font(bold=True)
    fills = {k: PatternFill("solid", fgColor=v) for k, v in list(FILLS.items()) + [("p-" + k, v) for k, v in PRIO_FILLS.items()]}
    white = PatternFill("solid", fgColor="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    from .report import safe_cell
    for n, r in enumerate(rows, start=2):          # 自己記列號：d.max_row 每次都要掃全表，幾萬列會很慢
        d.append([safe_cell(r[c]) for c in COLUMNS])     # OSV 說明是外部資料：防公式注入
        d.cell(n, 2).fill = fills.get("p-" + (r["priority"] or ""), white)
        for i, (col, _) in enumerate(RELEASES):
            d.cell(n, 5 + i).fill = fills.get(r["codes"][col], white)
        d.cell(n, 8).alignment = wrap
    widths = {"A": 17, "B": 10, "C": 11, "D": 22, "E": 30, "F": 30, "G": 30, "H": 80}
    for k, w in widths.items():
        d.column_dimensions[k].width = w
    d.freeze_panes = "A2"
    d.auto_filter.ref = d.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_executive(ws, summary, date_from, date_to, osv_date, vex_date):
    from openpyxl.styles import Alignment, Font, PatternFill
    e = summary["executive"]
    ws.append(["主管摘要（%s～%s，CVE 公開日期；狀態截至 %s）" % (date_from, date_to, date_to)])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.append([e["headline"]])
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
    ws.row_dimensions[2].height = 60
    ws.append([])
    ws.append(["%s（%d 個版本合計，同一 CVE 在不同版本分開計）" % (e["top_label"], len(RELEASES)), "適用", "官方已有修正（可處理）",
               "期末時尚未修正", "其中期末後已釋出修正", "不受影響"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append(["適用＝已有修正＋尚未修正＋不受影響", e["cards"]["applicable"], e["cards"]["fixed"], e["cards"]["open"], e["cards"]["later"],
               e["cards"]["not_affected"]])
    ws.append([])
    ws.append(["尚未解決／適用此版本" + ("（▲▼＝與上一期 %s～%s 比較）" % summary["prev_range"] if summary.get("prev_range") else "")])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["Ubuntu"] + summary["levels"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for m in e["matrix"]:
        ws.append([m["release"]] + [_matrix_text(c) for c in m["cells"]])
        for i, c in enumerate(m["cells"]):
            if c["open"] and c["level"] in ("critical", "high"):
                ws.cell(ws.max_row, 2 + i).fill = PatternFill("solid", fgColor=FILLS["needs-evaluation" if c["level"] == "critical" else "vulnerable"])
    ws.append([])
    ws.append(["要做的事"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["Ubuntu", "處理方式", "可解決／涉及的 CVE 數（%s）" % "、".join(summary["levels"])])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for a in e["actions"]:
        ws.append([a["release"], a["what"], a["count"]])
    ws.append([])
    ws.append(["資料：OSV %s、VEX %s；細節見「統計」「Kernel 修補版本」「狀態明細」「明細」各頁" % (osv_date or "", vex_date or "")])
    for col, w in zip("ABCDEF", (46, 22, 26, 26, 14, 14)):
        ws.column_dimensions[col].width = w


def _matrix_text(c):
    t = "%d／%d" % (c["open"], c["applicable"])
    if c["delta"]:
        t += "（%s%d）" % ("▲" if c["delta"] > 0 else "▼", abs(c["delta"]))
    return t


def _xlsx_overview(ws, summary):
    from openpyxl.styles import Font, PatternFill
    ws.append(["版本 × 等級 × 類別（以 CVE 計；適用＝扣掉「版本無此套件」與「官方無紀錄」；已修正＝期末前；尚未解決含期末後才修正）"])
    ws.cell(1, 1).font = Font(bold=True)
    ws.append(["Ubuntu", "等級", "類別", "適用此版本", "不受影響", "已修正", "其中日期不明", "尚未解決"])
    for c in ws[2]:
        c.font = Font(bold=True)
    ov = summary["overview"]
    for col, rel in RELEASES:
        for lv in summary["levels"]:
            for kind, label in KINDS:
                c = ov[(col, lv, kind)]
                ws.append([rel, lv, label, c["applicable"], c["not_affected"], c["fixed"], c["nodate"], c["open"]])
                ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=PRIO_FILLS[lv])
    for col, w in zip("ABCDEFGH", (9, 10, 10, 12, 10, 10, 12, 10)):
        ws.column_dimensions[col].width = w


def filename(date_from, date_to):
    return "cve-quarterly-%s-%s.xlsx" % (date_from.replace("-", ""), date_to.replace("-", ""))


def valid_date(s):
    """只接受 YYYY-MM-DD（Python 3.11 起 fromisoformat 也接受 20260701、2026-W27-1，這些會讓字串比較出錯）"""
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", s or ""):
        raise ValueError("日期格式要是 YYYY-MM-DD")
    d = datetime.date.fromisoformat(s)
    if not 2000 <= d.year <= 2100:
        raise ValueError("日期超出範圍")
    return d.isoformat()
